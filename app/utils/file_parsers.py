"""上传文件 → 纯文本 的解析器集合（无状态、无副作用）。

- .txt/.md：直接读文本
- .docx：python-docx 提取段落与表格（表格行以 " | " 连接）
- .pptx：python-pptx 逐页提取文本框、表格与演讲者备注
- 图片（.png/.jpg/.jpeg/.webp/.gif）：视觉模型 caption，模型由调用方传入；
  调用失败时**异常原样抛出**，由上层转成用户可读文案（不能静默吞掉，
  否则无法区分「模型不支持视觉」与「图片本身无内容」）。

不用 openjiuwen 自带的 AutoFileParser：它会把 docx 内嵌图片抽取到 ./images/
目录（污染工作目录），且 caption 失败时静默返回空串。
"""

import base64
import mimetypes
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openjiuwen.core.foundation.llm.model import Model
from openjiuwen.core.retrieval.indexing.processor.parser.captioner import (
    IMAGE_CAPTION_PROMPT,
)
from pptx import Presentation

TEXT_EXTS = {".txt", ".md"}
OFFICE_EXTS = {".docx", ".pptx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
SUPPORTED_EXTS = TEXT_EXTS | OFFICE_EXTS | IMAGE_EXTS
# 老二进制格式解析库不支持，明确拒绝并提示用户转格式
LEGACY_EXTS = {".doc", ".ppt"}


def parse_text(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="replace")


def _table_rows_text(rows) -> str:
    lines = []
    for row in rows:
        cells = [c.strip().replace("\n", " ") for c in row]
        line = " | ".join(c for c in cells).strip(" |")
        if line:
            lines.append(line)
    return "\n".join(lines)


def parse_docx(path: str) -> str:
    doc = Document(path)
    parts = []
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                parts.append(text)
        elif isinstance(block, Table):
            text = _table_rows_text(
                [[cell.text for cell in row.cells] for row in block.rows]
            )
            if text:
                parts.append(text)
    return "\n".join(parts)


def parse_pptx(path: str) -> str:
    prs = Presentation(path)
    pages = []
    for index, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
            elif shape.has_table:
                text = _table_rows_text(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
                if text:
                    lines.append(text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"备注：{notes}")
        if lines:
            pages.append(f"--- 第 {index} 页 ---\n" + "\n".join(lines))
    return "\n\n".join(pages)


async def caption_image(path: str, model: Model) -> str:
    """用视觉模型给图片生成文字描述。失败抛异常（上层负责转友好文案）。"""
    mime = mimetypes.guess_type(path)[0] or "image/png"
    b64 = base64.b64encode(Path(path).read_bytes()).decode("utf-8")
    resp = await model.invoke(
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": IMAGE_CAPTION_PROMPT},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime};base64,{b64}"},
                    },
                ],
            }
        ]
    )
    text = resp.content if hasattr(resp, "content") else str(resp)
    return (text or "").strip()
